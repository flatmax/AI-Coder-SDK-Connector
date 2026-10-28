"""SVG → PowerPoint export as native, editable shapes.

The SVG viewer's "Export as PowerPoint" action sends the current
SVG here and downloads the result. Every element becomes a real
DrawingML object rather than a picture, so the slide's contents
can be selected, restyled, and copied into another deck —
including PowerPoint for the web, which cannot paste vector
objects from outside Office.

Mapping:

- ``rect`` / ``circle`` / ``ellipse`` → preset rectangle, rounded
  rectangle, or ellipse, rotated when the element's transform
  rotates. A transform that skews, or a rect with elliptical
  corners, falls back to a freeform so the outline stays exact.
- ``line`` and single straight segments → connector lines, so
  PowerPoint offers endpoint handles.
- ``path`` / ``polyline`` / ``polygon`` → freeform (``custGeom``)
  in slide coordinates; arcs and quadratics become cubics.
- ``text`` / ``tspan`` → text boxes, one per label, with runs
  carrying each tspan's font, weight, style, and colour.
- ``g`` / ``a`` / nested ``svg`` / ``use`` → group shapes.
- ``image`` → picture, from a data URI or a repo-relative href.
- Fill and stroke colours, opacity, linear and radial
  gradients, dash arrays, caps, joins, and arrow markers map to
  their DrawingML equivalents. ``<style>`` sheets are honoured
  for tag, class, and id selectors, compound and descendant.

Clip paths, masks, filters, and patterns have no DrawingML
counterpart; they are dropped and reported in the warnings the
caller shows the user, so nothing disappears silently.

One SVG user unit at the SVG's authored size is one CSS pixel
(9525 EMU), so the slide is the size a browser draws the SVG.
"""

from __future__ import annotations

import base64
import colorsys
import io
import math
import re
import xml.etree.ElementTree as ET
from collections.abc import Callable
from dataclasses import dataclass, field
from xml.sax.saxutils import escape

from aic_dc.doc_index.extractors.svg_geometry import (
    IDENTITY,
    Matrix,
    compose,
    parse_transform,
    transform_point,
)

EMU_PER_PX = 9525
_EMU_PER_PT = 12700
# PowerPoint's slide-size limits.
_MIN_SLIDE_EMU = 914400
_MAX_SLIDE_EMU = 51206400
# PowerPoint draws single-spaced text at about 1.2 em per line,
# with the baseline about a quarter em above the line's bottom.
_LINE_PITCH_EM = 1.2
_DESCENT_EM = 0.25
# Average advance of a character, for sizing text boxes. The box
# never wraps, so this only affects the selection frame.
_CHAR_WIDTH_EM = 0.55

_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_NS_R = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
_NSDECLS = f'xmlns:a="{_NS_A}" xmlns:p="{_NS_P}" xmlns:r="{_NS_R}"'


class SvgExportError(ValueError):
    """The SVG could not be parsed or holds nothing to export."""


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------


def svg_to_pptx(
    svg_text: str,
    resolve_href: Callable[[str], bytes | None] | None = None,
) -> tuple[bytes, list[str]]:
    """Convert ``svg_text`` to a one-slide .pptx.

    ``resolve_href`` returns the bytes behind an ``<image>``
    href that isn't a data URI (a path relative to the SVG), or
    None when it can't. Returns the .pptx bytes and the
    human-readable warnings for anything approximated or
    dropped. Raises :class:`SvgExportError` on unparseable
    input.
    """
    from pptx import Presentation
    from pptx.util import Emu

    try:
        root = ET.fromstring(svg_text)
    except ET.ParseError as exc:
        raise SvgExportError(f"Not a valid SVG: {exc}") from exc
    _strip_namespaces(root)
    if root.tag != "svg":
        raise SvgExportError("The document's root element isn't <svg>")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    converter = _Converter(root, slide.part, resolve_href)
    nodes, width_emu, height_emu = converter.convert()
    if not nodes:
        raise SvgExportError("The SVG has nothing to export")

    prs.slide_width = Emu(width_emu)
    prs.slide_height = Emu(height_emu)
    # The template's size is named 4:3; this one is custom.
    sld_sz = prs.part._element.find(f"{{{_NS_P}}}sldSz")
    if sld_sz is not None:
        sld_sz.attrib.pop("type", None)
    sp_tree = slide.shapes._spTree
    ids = iter(range(2, 1 << 30))
    from pptx.oxml import parse_xml

    for node in nodes:
        sp_tree.append(parse_xml(node.to_xml(ids)))

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), converter.warnings()


# ---------------------------------------------------------------------------
# Output nodes
# ---------------------------------------------------------------------------


@dataclass
class _Box:
    x: float
    y: float
    w: float
    h: float

    def union(self, other: _Box) -> _Box:
        x0 = min(self.x, other.x)
        y0 = min(self.y, other.y)
        x1 = max(self.x + self.w, other.x + other.w)
        y1 = max(self.y + self.h, other.y + other.h)
        return _Box(x0, y0, x1 - x0, y1 - y0)


def _xfrm(box: _Box, rot: float = 0.0, flip_h=False, flip_v=False,
          tag: str = "a:xfrm") -> str:
    attrs = ""
    rot_val = round(rot * 60000) % 21600000
    if rot_val:
        attrs += f' rot="{rot_val}"'
    if flip_h:
        attrs += ' flipH="1"'
    if flip_v:
        attrs += ' flipV="1"'
    return (
        f"<{tag}{attrs}>"
        f'<a:off x="{round(box.x)}" y="{round(box.y)}"/>'
        f'<a:ext cx="{max(round(box.w), 0)}" cy="{max(round(box.h), 0)}"/>'
        f"</{tag}>"
    )


@dataclass
class _Shape:
    """One leaf object; ``template`` has ``{id}`` and ``{name}`` slots."""

    box: _Box
    template: str
    name: str

    def to_xml(self, ids) -> str:
        return self.template.format(
            id=next(ids), name=escape(self.name, {'"': "&quot;"}),
        )


@dataclass
class _Group:
    children: list
    name: str

    @property
    def box(self) -> _Box:
        box = self.children[0].box
        for child in self.children[1:]:
            box = box.union(child.box)
        return box

    def to_xml(self, ids) -> str:
        gid = next(ids)
        box = self.box
        xfrm = _xfrm(box).replace("</a:xfrm>", "")
        xfrm += (
            f'<a:chOff x="{round(box.x)}" y="{round(box.y)}"/>'
            f'<a:chExt cx="{round(box.w)}" cy="{round(box.h)}"/></a:xfrm>'
        )
        inner = "".join(child.to_xml(ids) for child in self.children)
        name = escape(self.name, {'"': "&quot;"})
        return (
            f"<p:grpSp {_NSDECLS}>"
            f'<p:nvGrpSpPr><p:cNvPr id="{gid}" name="{name}"/>'
            "<p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>"
            f"<p:grpSpPr>{xfrm}</p:grpSpPr>"
            f"{inner}</p:grpSp>"
        )


# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------


_INHERITED = {
    "fill", "fill-opacity", "stroke", "stroke-width", "stroke-opacity",
    "stroke-dasharray", "stroke-linecap", "stroke-linejoin",
    "font-family", "font-size", "font-weight", "font-style",
    "text-anchor", "dominant-baseline", "text-decoration",
    "visibility", "color", "marker-start", "marker-end",
}
_NOT_INHERITED = {
    "opacity", "display", "stop-color", "stop-opacity", "clip-path",
    "mask", "filter", "alignment-baseline",
}
_PROPS = _INHERITED | _NOT_INHERITED | {"marker"}

_INITIAL = {
    "fill": "black",
    "stroke": "none",
    "stroke-width": "1",
    "font-family": "",
    "font-size": 16.0,
    "text-anchor": "start",
    "color": "black",
}


@dataclass
class _Rule:
    chain: list  # [(combinator, (tag, ids, classes)), ...] left to right
    specificity: tuple
    order: int
    decls: dict


_COMPOUND_RE = re.compile(r"^(\*|[A-Za-z][\w-]*)?((?:[.#][\w-]+)*)$")


def _parse_declarations(text: str) -> dict[str, str]:
    decls = {}
    for part in text.split(";"):
        if ":" not in part:
            continue
        name, value = part.split(":", 1)
        name = name.strip().lower()
        value = value.replace("!important", "").strip()
        if name and value:
            decls[name] = value
    return decls


def _parse_stylesheet(css: str) -> list[_Rule]:
    css = re.sub(r"/\*.*?\*/", "", css, flags=re.S)
    rules = []
    for match in re.finditer(r"([^{}]+)\{([^{}]*)\}", css):
        decls = _parse_declarations(match.group(2))
        if not decls:
            continue
        for selector in match.group(1).split(","):
            chain = _parse_selector(selector.strip())
            if chain is None:
                continue
            ids = sum(len(c[1][1]) for c in chain)
            classes = sum(len(c[1][2]) for c in chain)
            tags = sum(1 for c in chain if c[1][0])
            rules.append(_Rule(chain, (ids, classes, tags), len(rules), decls))
    return rules


def _parse_selector(selector: str):
    """Parse ``tag.cls#id`` compounds joined by space or ``>``."""
    if not selector or selector.startswith("@"):
        return None
    tokens = re.split(r"\s*(>)\s*|\s+", selector)
    chain = []
    combinator = " "
    for token in tokens:
        if token is None or token == "":
            continue
        if token == ">":
            combinator = ">"
            continue
        match = _COMPOUND_RE.match(token)
        if match is None or not token:
            return None
        tag = match.group(1)
        tag = None if tag in (None, "*") else tag
        rest = match.group(2) or ""
        ids = re.findall(r"#([\w-]+)", rest)
        classes = re.findall(r"\.([\w-]+)", rest)
        chain.append((combinator, (tag, ids, classes)))
        combinator = " "
    return chain or None


def _compound_matches(compound, el: ET.Element) -> bool:
    tag, ids, classes = compound
    if tag and el.tag != tag:
        return False
    if ids and el.get("id") not in ids:
        return False
    if classes:
        have = set((el.get("class") or "").split())
        if not set(classes) <= have:
            return False
    return True


def _rule_matches(rule: _Rule, el: ET.Element, ancestors: list) -> bool:
    chain = rule.chain
    if not _compound_matches(chain[-1][1], el):
        return False
    pos = len(ancestors)
    for i in range(len(chain) - 1, 0, -1):
        combinator = chain[i][0]
        compound = chain[i - 1][1]
        if combinator == ">":
            pos -= 1
            if pos < 0 or not _compound_matches(compound, ancestors[pos]):
                return False
        else:
            pos -= 1
            while pos >= 0 and not _compound_matches(compound, ancestors[pos]):
                pos -= 1
            if pos < 0:
                return False
    return True


# ---------------------------------------------------------------------------
# Values
# ---------------------------------------------------------------------------


_NUM_RE = re.compile(r"[-+]?(?:\d+\.\d*|\.\d+|\d+)(?:[eE][-+]?\d+)?")
_UNIT_PX = {
    "": 1.0, "px": 1.0, "pt": 96 / 72, "pc": 16.0,
    "mm": 96 / 25.4, "cm": 96 / 2.54, "in": 96.0, "q": 96 / 101.6,
}


def _length(value, font_size: float = 16.0, ref: float = 0.0,
            default: float | None = None) -> float | None:
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return float(value)
    match = re.match(
        r"\s*([-+]?(?:\d+\.\d*|\.\d+|\d+)(?:[eE][-+]?\d+)?)\s*([a-zA-Z%]*)",
        value,
    )
    if not match:
        return default
    number = float(match.group(1))
    unit = match.group(2).lower()
    if unit in _UNIT_PX:
        return number * _UNIT_PX[unit]
    if unit == "em":
        return number * font_size
    if unit == "ex":
        return number * font_size / 2
    if unit == "%":
        return number / 100 * ref
    return number


def _first_length(value, font_size=16.0, ref=0.0, default=0.0) -> float:
    """First entry of a per-glyph list such as ``x="10 20 30"``."""
    if value is None:
        return default
    first = re.split(r"[\s,]+", value.strip())[0] if value.strip() else ""
    return _length(first, font_size, ref, default)


def _numbers(text: str | None) -> list[float]:
    return [float(m) for m in _NUM_RE.findall(text or "")]


_NAMED = dict(item.split(":") for item in (
    "aliceblue:f0f8ff antiquewhite:faebd7 aqua:00ffff aquamarine:7fffd4 "
    "azure:f0ffff beige:f5f5dc bisque:ffe4c4 black:000000 "
    "blanchedalmond:ffebcd blue:0000ff blueviolet:8a2be2 brown:a52a2a "
    "burlywood:deb887 cadetblue:5f9ea0 chartreuse:7fff00 chocolate:d2691e "
    "coral:ff7f50 cornflowerblue:6495ed cornsilk:fff8dc crimson:dc143c "
    "cyan:00ffff darkblue:00008b darkcyan:008b8b darkgoldenrod:b8860b "
    "darkgray:a9a9a9 darkgreen:006400 darkgrey:a9a9a9 darkkhaki:bdb76b "
    "darkmagenta:8b008b darkolivegreen:556b2f darkorange:ff8c00 "
    "darkorchid:9932cc darkred:8b0000 darksalmon:e9967a "
    "darkseagreen:8fbc8f darkslateblue:483d8b darkslategray:2f4f4f "
    "darkslategrey:2f4f4f darkturquoise:00ced1 darkviolet:9400d3 "
    "deeppink:ff1493 deepskyblue:00bfff dimgray:696969 dimgrey:696969 "
    "dodgerblue:1e90ff firebrick:b22222 floralwhite:fffaf0 "
    "forestgreen:228b22 fuchsia:ff00ff gainsboro:dcdcdc "
    "ghostwhite:f8f8ff gold:ffd700 goldenrod:daa520 gray:808080 "
    "green:008000 greenyellow:adff2f grey:808080 honeydew:f0fff0 "
    "hotpink:ff69b4 indianred:cd5c5c indigo:4b0082 ivory:fffff0 "
    "khaki:f0e68c lavender:e6e6fa lavenderblush:fff0f5 lawngreen:7cfc00 "
    "lemonchiffon:fffacd lightblue:add8e6 lightcoral:f08080 "
    "lightcyan:e0ffff lightgoldenrodyellow:fafad2 lightgray:d3d3d3 "
    "lightgreen:90ee90 lightgrey:d3d3d3 lightpink:ffb6c1 "
    "lightsalmon:ffa07a lightseagreen:20b2aa lightskyblue:87cefa "
    "lightslategray:778899 lightslategrey:778899 lightsteelblue:b0c4de "
    "lightyellow:ffffe0 lime:00ff00 limegreen:32cd32 linen:faf0e6 "
    "magenta:ff00ff maroon:800000 mediumaquamarine:66cdaa "
    "mediumblue:0000cd mediumorchid:ba55d3 mediumpurple:9370db "
    "mediumseagreen:3cb371 mediumslateblue:7b68ee "
    "mediumspringgreen:00fa9a mediumturquoise:48d1cc "
    "mediumvioletred:c71585 midnightblue:191970 mintcream:f5fffa "
    "mistyrose:ffe4e1 moccasin:ffe4b5 navajowhite:ffdead navy:000080 "
    "oldlace:fdf5e6 olive:808000 olivedrab:6b8e23 orange:ffa500 "
    "orangered:ff4500 orchid:da70d6 palegoldenrod:eee8aa "
    "palegreen:98fb98 paleturquoise:afeeee palevioletred:db7093 "
    "papayawhip:ffefd5 peachpuff:ffdab9 peru:cd853f pink:ffc0cb "
    "plum:dda0dd powderblue:b0e0e6 purple:800080 rebeccapurple:663399 "
    "red:ff0000 rosybrown:bc8f8f royalblue:4169e1 saddlebrown:8b4513 "
    "salmon:fa8072 sandybrown:f4a460 seagreen:2e8b57 seashell:fff5ee "
    "sienna:a0522d silver:c0c0c0 skyblue:87ceeb slateblue:6a5acd "
    "slategray:708090 slategrey:708090 snow:fffafa springgreen:00ff7f "
    "steelblue:4682b4 tan:d2b48c teal:008080 thistle:d8bfd8 "
    "tomato:ff6347 turquoise:40e0d0 violet:ee82ee wheat:f5deb3 "
    "white:ffffff whitesmoke:f5f5f5 yellow:ffff00 yellowgreen:9acd32"
).split())


def _parse_color(value: str | None, current: str = "black"):
    """Return ``(RRGGBB, alpha)``, ``("url", id)``, or None for none."""
    if value is None:
        return None
    value = value.strip()
    low = value.lower()
    if low.startswith("url("):
        match = re.match(r"url\(\s*['\"]?#([^)'\"]+)['\"]?\s*\)\s*(.*)", value)
        if match:
            return ("url", match.group(1), match.group(2).strip())
        return None
    if low in ("none", "transparent", ""):
        return None
    if low == "currentcolor":
        return _parse_color(current) if current.lower() != "currentcolor" else None
    if low in _NAMED:
        return (_NAMED[low].upper(), 1.0)
    if low.startswith("#"):
        digits = low[1:]
        if len(digits) in (3, 4):
            digits = "".join(c * 2 for c in digits)
        if len(digits) == 6:
            return (digits.upper(), 1.0)
        if len(digits) == 8:
            return (digits[:6].upper(), int(digits[6:], 16) / 255)
        return None
    match = re.match(r"(rgba?|hsla?)\(([^)]*)\)", low)
    if match:
        parts = [p for p in re.split(r"[\s,/]+", match.group(2)) if p]
        if len(parts) < 3:
            return None
        alpha = 1.0
        if len(parts) >= 4:
            alpha = _length(parts[3], ref=1.0, default=1.0)
        if match.group(1).startswith("rgb"):
            channels = [
                _length(p, ref=255.0, default=0.0) for p in parts[:3]
            ]
        else:
            hue = (_numbers(parts[0]) or [0])[0] / 360
            sat = _length(parts[1], ref=1.0, default=0.0)
            lig = _length(parts[2], ref=1.0, default=0.0)
            r, g, b = colorsys.hls_to_rgb(hue % 1, lig, sat)
            channels = [r * 255, g * 255, b * 255]
        hexa = "".join(
            f"{max(0, min(255, round(c))):02X}" for c in channels
        )
        return (hexa, max(0.0, min(1.0, alpha)))
    return None


def _opacity(value, default=1.0) -> float:
    number = _length(value, ref=1.0, default=default)
    return max(0.0, min(1.0, number))


def _color_xml(hexa: str, alpha: float) -> str:
    if alpha >= 0.999:
        return f'<a:srgbClr val="{hexa}"/>'
    return (
        f'<a:srgbClr val="{hexa}"><a:alpha val="{round(alpha * 100000)}"/>'
        "</a:srgbClr>"
    )


# ---------------------------------------------------------------------------
# Matrix helpers
# ---------------------------------------------------------------------------


def _decompose(m: Matrix):
    """Return ``(sx, sy, degrees, skewed)`` for ``m``'s linear part."""
    sx = math.hypot(m.a, m.b)
    if sx < 1e-12:
        return 0.0, 0.0, 0.0, True
    det = m.a * m.d - m.b * m.c
    sy = det / sx
    shear = (m.a * m.c + m.b * m.d) / sx
    angle = math.degrees(math.atan2(m.b, m.a))
    skewed = abs(shear) > 1e-4 * max(sx, abs(sy)) or sy <= 0
    return sx, abs(sy), angle, skewed


def _parse_transform(text: str | None) -> Matrix:
    """``parse_transform`` plus ``skewX`` / ``skewY``, which it skips."""
    if not text or "skew" not in text:
        return parse_transform(text)
    result = IDENTITY
    for match in re.finditer(r"([a-zA-Z]+)\s*\(([^)]*)\)", text):
        name = match.group(1)
        args = _numbers(match.group(2))
        if name in ("skewX", "skewY") and args:
            t = math.tan(math.radians(args[0]))
            step = Matrix(1, 0, t, 1, 0, 0) if name == "skewX" else (
                Matrix(1, t, 0, 1, 0, 0))
        else:
            step = parse_transform(match.group(0))
        result = compose(result, step)
    return result


def _mean_scale(m: Matrix) -> float:
    return math.sqrt(abs(m.a * m.d - m.b * m.c))


def _translate(x: float, y: float) -> Matrix:
    return Matrix(1, 0, 0, 1, x, y)


def _viewbox_matrix(viewbox, width, height, par: str | None) -> Matrix:
    """Map a viewBox into a ``width`` × ``height`` viewport."""
    vx, vy, vw, vh = viewbox
    sx, sy = width / vw, height / vh
    par = (par or "xMidYMid meet").split()
    align = par[0] if par else "xMidYMid"
    if align == "none":
        return Matrix(sx, 0, 0, sy, -vx * sx, -vy * sy)
    scale = max(sx, sy) if "slice" in par else min(sx, sy)
    tx = -vx * scale
    ty = -vy * scale
    if "xMid" in align:
        tx += (width - vw * scale) / 2
    elif "xMax" in align:
        tx += width - vw * scale
    if "YMid" in align:
        ty += (height - vh * scale) / 2
    elif "YMax" in align:
        ty += height - vh * scale
    return Matrix(scale, 0, 0, scale, tx, ty)


def _parse_viewbox(value: str | None):
    numbers = _numbers(value)
    if len(numbers) == 4 and numbers[2] > 0 and numbers[3] > 0:
        return tuple(numbers)
    return None


# ---------------------------------------------------------------------------
# Path data
# ---------------------------------------------------------------------------


def _tokenize_path(d: str):
    """Yield command letters and numbers; arc flags may be unseparated."""
    tokens = []
    i = 0
    n = len(d)
    command = None
    arg_index = 0
    while i < n:
        ch = d[i]
        if ch.isalpha() and ch not in "eE":
            command = ch
            arg_index = 0
            tokens.append(ch)
            i += 1
            continue
        if ch in " \t\r\n,":
            i += 1
            continue
        if command in ("a", "A") and arg_index % 7 in (3, 4) and ch in "01":
            tokens.append(float(ch))
            arg_index += 1
            i += 1
            continue
        match = _NUM_RE.match(d, i)
        if not match:
            i += 1
            continue
        tokens.append(float(match.group(0)))
        arg_index += 1
        i = match.end()
    return tokens


def _arc_to_cubics(x0, y0, rx, ry, phi_deg, large, sweep, x, y):
    """Endpoint-parameterised arc → cubic Béziers (SVG spec F.6)."""
    if (x0, y0) == (x, y):
        return []
    rx, ry = abs(rx), abs(ry)
    if rx == 0 or ry == 0:
        return [(x0, y0, x, y, x, y)]
    phi = math.radians(phi_deg)
    cos_p, sin_p = math.cos(phi), math.sin(phi)
    dx, dy = (x0 - x) / 2, (y0 - y) / 2
    x1p = cos_p * dx + sin_p * dy
    y1p = -sin_p * dx + cos_p * dy
    lam = (x1p / rx) ** 2 + (y1p / ry) ** 2
    if lam > 1:
        rx *= math.sqrt(lam)
        ry *= math.sqrt(lam)
    num = rx * rx * ry * ry - rx * rx * y1p * y1p - ry * ry * x1p * x1p
    den = rx * rx * y1p * y1p + ry * ry * x1p * x1p
    coef = math.sqrt(max(0.0, num / den)) if den else 0.0
    if large == sweep:
        coef = -coef
    cxp = coef * rx * y1p / ry
    cyp = -coef * ry * x1p / rx
    cx = cos_p * cxp - sin_p * cyp + (x0 + x) / 2
    cy = sin_p * cxp + cos_p * cyp + (y0 + y) / 2

    def angle(ux, uy, vx, vy):
        return math.atan2(ux * vy - uy * vx, ux * vx + uy * vy)

    theta1 = angle(1, 0, (x1p - cxp) / rx, (y1p - cyp) / ry)
    delta = angle(
        (x1p - cxp) / rx, (y1p - cyp) / ry,
        (-x1p - cxp) / rx, (-y1p - cyp) / ry,
    )
    if not sweep and delta > 0:
        delta -= 2 * math.pi
    elif sweep and delta < 0:
        delta += 2 * math.pi
    segments = max(1, math.ceil(abs(delta) / (math.pi / 2) - 1e-9))
    step = delta / segments
    k = 4 / 3 * math.tan(step / 4)

    def point(t):
        return (
            cx + rx * math.cos(t) * cos_p - ry * math.sin(t) * sin_p,
            cy + rx * math.cos(t) * sin_p + ry * math.sin(t) * cos_p,
        )

    def deriv(t):
        return (
            -rx * math.sin(t) * cos_p - ry * math.cos(t) * sin_p,
            -rx * math.sin(t) * sin_p + ry * math.cos(t) * cos_p,
        )

    cubics = []
    t = theta1
    for _ in range(segments):
        p0, p3 = point(t), point(t + step)
        d0, d3 = deriv(t), deriv(t + step)
        cubics.append((
            p0[0] + k * d0[0], p0[1] + k * d0[1],
            p3[0] - k * d3[0], p3[1] - k * d3[1],
            p3[0], p3[1],
        ))
        t += step
    # Land exactly on the endpoint.
    last = cubics[-1]
    cubics[-1] = last[:4] + (x, y)
    return cubics


def _parse_path(d: str | None) -> list:
    """Absolute segments: ``("M", x, y)``, ``("L", x, y)``,
    ``("C", x1, y1, x2, y2, x, y)``, ``("Z",)``."""
    tokens = _tokenize_path(d or "")
    segments = []
    i = 0
    cx = cy = sx = sy = 0.0
    last_ctrl = None  # (kind, x, y) of the previous curve's control
    command = None

    def take(count):
        nonlocal i
        values = tokens[i:i + count]
        if len(values) < count or any(isinstance(v, str) for v in values):
            raise IndexError
        i += count
        return values

    while i < len(tokens):
        token = tokens[i]
        if isinstance(token, str):
            command = token
            i += 1
            if command in "zZ":
                segments.append(("Z",))
                cx, cy = sx, sy
                last_ctrl = None
                continue
        elif command is None:
            break
        rel = command.islower()
        op = command.upper()
        try:
            if op == "M":
                x, y = take(2)
                if rel:
                    x, y = cx + x, cy + y
                segments.append(("M", x, y))
                cx, cy = sx, sy = x, y
                command = "l" if rel else "L"
                last_ctrl = None
            elif op in ("L", "H", "V"):
                if op == "L":
                    x, y = take(2)
                    if rel:
                        x, y = cx + x, cy + y
                elif op == "H":
                    (x,) = take(1)
                    x, y = (cx + x if rel else x), cy
                else:
                    (y,) = take(1)
                    x, y = cx, (cy + y if rel else y)
                segments.append(("L", x, y))
                cx, cy = x, y
                last_ctrl = None
            elif op in ("C", "S"):
                if op == "C":
                    x1, y1, x2, y2, x, y = take(6)
                    if rel:
                        x1, y1 = cx + x1, cy + y1
                else:
                    x2, y2, x, y = take(4)
                    if last_ctrl and last_ctrl[0] == "C":
                        x1, y1 = 2 * cx - last_ctrl[1], 2 * cy - last_ctrl[2]
                    else:
                        x1, y1 = cx, cy
                if rel:
                    x2, y2, x, y = cx + x2, cy + y2, cx + x, cy + y
                if op == "S" and rel and not (last_ctrl and last_ctrl[0] == "C"):
                    x1, y1 = cx, cy
                segments.append(("C", x1, y1, x2, y2, x, y))
                last_ctrl = ("C", x2, y2)
                cx, cy = x, y
            elif op in ("Q", "T"):
                if op == "Q":
                    qx, qy, x, y = take(4)
                    if rel:
                        qx, qy = cx + qx, cy + qy
                else:
                    x, y = take(2)
                    if last_ctrl and last_ctrl[0] == "Q":
                        qx, qy = 2 * cx - last_ctrl[1], 2 * cy - last_ctrl[2]
                    else:
                        qx, qy = cx, cy
                if rel:
                    x, y = cx + x, cy + y
                segments.append((
                    "C",
                    cx + 2 / 3 * (qx - cx), cy + 2 / 3 * (qy - cy),
                    x + 2 / 3 * (qx - x), y + 2 / 3 * (qy - y),
                    x, y,
                ))
                last_ctrl = ("Q", qx, qy)
                cx, cy = x, y
            elif op == "A":
                rx, ry, phi, large, sweep, x, y = take(7)
                if rel:
                    x, y = cx + x, cy + y
                for cubic in _arc_to_cubics(
                    cx, cy, rx, ry, phi, bool(large), bool(sweep), x, y,
                ):
                    segments.append(("C",) + cubic)
                cx, cy = x, y
                last_ctrl = None
            else:
                i += 1
        except IndexError:
            break
    if segments and segments[0][0] != "M":
        return []
    return segments


def _cubic_extrema(p0, p1, p2, p3):
    """Parameter values in (0, 1) where a 1-D cubic turns."""
    a = -p0 + 3 * p1 - 3 * p2 + p3
    b = 2 * (p0 - 2 * p1 + p2)
    c = p1 - p0
    roots = []
    if abs(a) < 1e-12:
        if abs(b) > 1e-12:
            roots.append(-c / b)
    else:
        disc = b * b - 4 * a * c
        if disc >= 0:
            sq = math.sqrt(disc)
            roots += [(-b + sq) / (2 * a), (-b - sq) / (2 * a)]
    return [t for t in roots if 0 < t < 1]


def _segments_bbox(segments) -> _Box | None:
    xs, ys = [], []
    px = py = 0.0
    for seg in segments:
        if seg[0] in ("M", "L"):
            px, py = seg[1], seg[2]
            xs.append(px)
            ys.append(py)
        elif seg[0] == "C":
            _, x1, y1, x2, y2, x, y = seg
            xs.append(x)
            ys.append(y)
            for values, out in (((px, x1, x2, x), xs), ((py, y1, y2, y), ys)):
                for t in _cubic_extrema(*values):
                    mt = 1 - t
                    out.append(
                        mt ** 3 * values[0] + 3 * mt * mt * t * values[1]
                        + 3 * mt * t * t * values[2] + t ** 3 * values[3]
                    )
            px, py = x, y
    if not xs:
        return None
    return _Box(min(xs), min(ys), max(xs) - min(xs), max(ys) - min(ys))


def _transform_segments(m: Matrix, segments) -> list:
    out = []
    for seg in segments:
        if seg[0] == "Z":
            out.append(seg)
            continue
        coords = []
        for j in range(1, len(seg), 2):
            coords += transform_point(m, seg[j], seg[j + 1])
        out.append((seg[0],) + tuple(coords))
    return out


# ---------------------------------------------------------------------------
# The converter
# ---------------------------------------------------------------------------


_SKIP_TAGS = {
    "defs", "title", "desc", "metadata", "style", "script", "marker",
    "symbol", "clipPath", "mask", "pattern", "linearGradient",
    "radialGradient", "filter",
}
_CONTAINERS = {"g", "a", "switch"}
_GENERIC_FONTS = {
    "serif": "Times New Roman",
    "sans-serif": "Arial",
    "monospace": "Courier New",
    "cursive": "Comic Sans MS",
    "fantasy": "Impact",
    "system-ui": "Segoe UI",
    "ui-sans-serif": "Arial",
    "ui-serif": "Times New Roman",
    "ui-monospace": "Courier New",
    "-apple-system": None,
    "blinkmacsystemfont": None,
}


@dataclass
class _Ctx:
    matrix: Matrix
    style: dict
    ancestors: list
    opacity: float = 1.0
    viewport: tuple = (0.0, 0.0)
    use_depth: int = 0


@dataclass
class _Run:
    text: str
    style: dict


@dataclass
class _Line:
    x: float
    y: float
    runs: list = field(default_factory=list)


class _Converter:
    def __init__(self, root: ET.Element, slide_part, resolve_href) -> None:
        self._root = root
        self._slide_part = slide_part
        self._resolve_href = resolve_href
        self._ids = {
            el.get("id"): el for el in root.iter() if el.get("id")
        }
        css = "\n".join(
            "".join(el.itertext()) for el in root.iter("style")
        )
        self._rules = _parse_stylesheet(css)
        self._warnings: dict[str, int] = {}
        self._counts: dict[str, int] = {}

    # -- bookkeeping ------------------------------------------------------

    def _warn(self, message: str) -> None:
        self._warnings[message] = self._warnings.get(message, 0) + 1

    def warnings(self) -> list[str]:
        return [
            message if count == 1 else f"{message} ({count}×)"
            for message, count in self._warnings.items()
        ]

    def _name(self, el: ET.Element, fallback: str) -> str:
        if el.get("id"):
            return el.get("id")
        self._counts[fallback] = self._counts.get(fallback, 0) + 1
        return f"{fallback} {self._counts[fallback]}"

    # -- top level --------------------------------------------------------

    def convert(self):
        root = self._root
        viewbox = _parse_viewbox(root.get("viewBox"))
        width = _length(root.get("width"), default=None)
        height = _length(root.get("height"), default=None)
        if root.get("width", "").strip().endswith("%"):
            width = None
        if root.get("height", "").strip().endswith("%"):
            height = None
        if viewbox:
            if width is None and height is None:
                width, height = viewbox[2], viewbox[3]
            elif width is None:
                width = height * viewbox[2] / viewbox[3]
            elif height is None:
                height = width * viewbox[3] / viewbox[2]
            to_px = _viewbox_matrix(
                viewbox, width, height, root.get("preserveAspectRatio"),
            )
            viewport = (viewbox[2], viewbox[3])
        else:
            to_px = IDENTITY
            viewport = (width or 0.0, height or 0.0)

        scale = EMU_PER_PX
        if width and height:
            largest = max(width, height) * EMU_PER_PX
            if largest > _MAX_SLIDE_EMU:
                scale = _MAX_SLIDE_EMU * EMU_PER_PX / largest
        matrix = compose(Matrix(scale, 0, 0, scale, 0, 0), to_px)
        style = self._computed(root, dict(_INITIAL), [])
        ctx = _Ctx(matrix, style, [root], viewport=viewport)
        nodes = self._children(root, ctx)

        if width and height:
            slide_w, slide_h = width * scale, height * scale
        else:
            # No authored size — fit the slide to the content.
            box = _Group(nodes, "").box if nodes else _Box(0, 0, 0, 0)
            slide_w = box.x + box.w
            slide_h = box.y + box.h
        def clamp(v):
            return round(min(max(v, _MIN_SLIDE_EMU), _MAX_SLIDE_EMU))

        return nodes, clamp(slide_w), clamp(slide_h)

    # -- style ------------------------------------------------------------

    def _computed(self, el, parent: dict, ancestors: list) -> dict:
        style = {k: v for k, v in parent.items() if k in _INHERITED}
        specified = {}
        for name in _PROPS:
            if el.get(name) is not None:
                specified[name] = el.get(name).strip()
        matched = [r for r in self._rules if _rule_matches(r, el, ancestors)]
        matched.sort(key=lambda r: (r.specificity, r.order))
        for rule in matched:
            specified.update(rule.decls)
        specified.update(_parse_declarations(el.get("style") or ""))
        if "marker" in specified:
            specified.setdefault("marker-start", specified["marker"])
            specified.setdefault("marker-end", specified["marker"])
        for name, value in specified.items():
            if name not in _PROPS or name == "marker":
                continue
            if value == "inherit":
                if name in parent:
                    style[name] = parent[name]
                continue
            if name == "font-size":
                base = parent.get("font-size", 16.0)
                named = {
                    "smaller": base / 1.2, "larger": base * 1.2,
                    "small": 13.0, "medium": 16.0, "large": 18.0,
                }
                if value in named:
                    style[name] = named[value]
                else:
                    style[name] = _length(value, base, base, base)
                continue
            style[name] = value
        return style

    # -- tree walk --------------------------------------------------------

    def _child_ctx(self, el, ctx: _Ctx, extra: Matrix = IDENTITY):
        style = self._computed(el, ctx.style, ctx.ancestors)
        matrix = compose(ctx.matrix, _parse_transform(el.get("transform")))
        matrix = compose(matrix, extra)
        opacity = ctx.opacity * _opacity(style.get("opacity"))
        return _Ctx(
            matrix, style, ctx.ancestors + [el], opacity,
            ctx.viewport, ctx.use_depth,
        )

    def _children(self, el, ctx: _Ctx) -> list:
        nodes = []
        for child in el:
            if child.tag == "switch":
                child = next(iter(child), None)
                if child is None:
                    continue
            nodes.extend(self._element(child, ctx))
        return nodes

    def _element(self, el, ctx: _Ctx) -> list:
        tag = el.tag
        if tag in _SKIP_TAGS or not isinstance(tag, str):
            return []
        local = self._child_ctx(el, ctx)
        style = local.style
        if style.get("display") == "none":
            return []
        for prop, label in (("clip-path", "Clip paths"),
                            ("mask", "Masks"), ("filter", "Filters")):
            value = style.get(prop)
            if value and value != "none":
                self._warn(f"{label} were ignored")

        if tag in _CONTAINERS:
            return self._group(el, self._children(el, local), "Group")
        if tag == "svg":
            return self._nested_svg(el, ctx)
        if tag == "use":
            return self._use(el, ctx)
        if style.get("visibility") in ("hidden", "collapse"):
            return []
        if tag == "text":
            return self._text(el, local)
        if tag == "image":
            return self._image(el, local)
        if tag == "foreignObject":
            self._warn("Embedded HTML (foreignObject) was skipped")
            return []
        return self._shape(el, local)

    def _group(self, el, nodes: list, label: str) -> list:
        if len(nodes) <= 1:
            return nodes
        return [_Group(nodes, self._name(el, label))]

    def _nested_svg(self, el, ctx: _Ctx) -> list:
        fs = ctx.style.get("font-size", 16.0)
        x = _length(el.get("x"), fs, ctx.viewport[0], 0.0)
        y = _length(el.get("y"), fs, ctx.viewport[1], 0.0)
        width = _length(el.get("width"), fs, ctx.viewport[0], ctx.viewport[0])
        height = _length(el.get("height"), fs, ctx.viewport[1], ctx.viewport[1])
        extra = _translate(x, y)
        viewbox = _parse_viewbox(el.get("viewBox"))
        viewport = (width, height)
        if viewbox and width and height:
            extra = compose(extra, _viewbox_matrix(
                viewbox, width, height, el.get("preserveAspectRatio"),
            ))
            viewport = (viewbox[2], viewbox[3])
        local = self._child_ctx(el, ctx, extra)
        local.viewport = viewport
        return self._group(el, self._children(el, local), "Group")

    def _use(self, el, ctx: _Ctx) -> list:
        href = el.get("href") or ""
        target = self._ids.get(href[1:]) if href.startswith("#") else None
        if target is None or ctx.use_depth > 8:
            return []
        fs = ctx.style.get("font-size", 16.0)
        x = _length(el.get("x"), fs, ctx.viewport[0], 0.0)
        y = _length(el.get("y"), fs, ctx.viewport[1], 0.0)
        local = self._child_ctx(el, ctx, _translate(x, y))
        local.use_depth += 1
        if local.style.get("display") == "none":
            return []
        if target.tag == "symbol":
            width = _length(el.get("width"), fs, ctx.viewport[0], None)
            height = _length(el.get("height"), fs, ctx.viewport[1], None)
            viewbox = _parse_viewbox(target.get("viewBox"))
            sym = self._child_ctx(target, local)
            if viewbox and width and height:
                sym.matrix = compose(sym.matrix, _viewbox_matrix(
                    viewbox, width, height,
                    target.get("preserveAspectRatio"),
                ))
            return self._group(el, self._children(target, sym), "Group")
        return self._group(el, self._element(target, local), "Group")

    # -- fills and strokes ------------------------------------------------

    def _gradient_xml(self, gid: str, extra_alpha: float, local_box,
                      matrix: Matrix, frame_rot: float) -> str | None:
        grad = self._ids.get(gid)
        if grad is None or grad.tag not in ("linearGradient", "radialGradient"):
            if grad is not None and grad.tag == "pattern":
                self._warn("Pattern fills were replaced with no fill")
            return None
        # Stops and attributes may come from an href'd gradient.
        chain = [grad]
        while len(chain) < 8:
            href = chain[-1].get("href") or ""
            nxt = self._ids.get(href[1:]) if href.startswith("#") else None
            if nxt is None or nxt in chain:
                break
            chain.append(nxt)

        def attr(name, default=None):
            for g in chain:
                if g.get(name) is not None:
                    return g.get(name)
            return default

        stop_owner = next((g for g in chain if g.findall("stop")), None)
        stops = []
        if stop_owner is not None:
            for stop in stop_owner.findall("stop"):
                st = self._computed(stop, {"color": "black"}, [stop_owner])
                color = _parse_color(st.get("stop-color", "black"),
                                     st.get("color", "black"))
                if color is None or color[0] == "url":
                    color = ("000000", 0.0)
                alpha = color[1] * _opacity(st.get("stop-opacity")) * extra_alpha
                offset = _length(stop.get("offset", "0"), ref=1.0, default=0.0)
                offset = max(0.0, min(1.0, offset))
                if stops and offset < stops[-1][0]:
                    offset = stops[-1][0]
                stops.append((offset, color[0], alpha))
        if not stops:
            return None
        if len(stops) == 1:
            return f"<a:solidFill>{_color_xml(stops[0][1], stops[0][2])}</a:solidFill>"
        gs = "".join(
            f'<a:gs pos="{round(o * 100000)}">{_color_xml(c, a)}</a:gs>'
            for o, c, a in stops
        )
        user_space = attr("gradientUnits") == "userSpaceOnUse"
        gt = _parse_transform(attr("gradientTransform"))
        if grad.tag == "linearGradient":
            fs = 16.0
            ref_w = local_box.w if not user_space else 1.0
            ref_h = local_box.h if not user_space else 1.0
            x1 = _length(attr("x1", "0%"), fs, 1.0 if not user_space else 0, 0.0)
            y1 = _length(attr("y1", "0%"), fs, 1.0 if not user_space else 0, 0.0)
            x2 = _length(attr("x2", "100%"), fs, 1.0 if not user_space else 0, 1.0)
            y2 = _length(attr("y2", "0%"), fs, 1.0 if not user_space else 0, 0.0)
            vx, vy = (x2 - x1) * ref_w, (y2 - y1) * ref_h
            vx, vy = gt.a * vx + gt.c * vy, gt.b * vx + gt.d * vy
            wx, wy = matrix.a * vx + matrix.c * vy, matrix.b * vx + matrix.d * vy
            angle = (math.degrees(math.atan2(wy, wx)) - frame_rot) % 360
            return (
                f'<a:gradFill rotWithShape="1"><a:gsLst>{gs}</a:gsLst>'
                f'<a:lin ang="{round(angle * 60000) % 21600000}" scaled="0"/>'
                "</a:gradFill>"
            )
        cx = _length(attr("cx", "50%"), ref=1.0, default=0.5)
        cy = _length(attr("cy", "50%"), ref=1.0, default=0.5)
        if user_space and local_box.w and local_box.h:
            cx = (cx - local_box.x) / local_box.w
            cy = (cy - local_box.y) / local_box.h
        cx, cy = max(0.0, min(1.0, cx)), max(0.0, min(1.0, cy))
        left, top = round(cx * 100000), round(cy * 100000)
        return (
            f'<a:gradFill rotWithShape="1"><a:gsLst>{gs}</a:gsLst>'
            f'<a:path path="circle"><a:fillToRect l="{left}" t="{top}" '
            f'r="{100000 - left}" b="{100000 - top}"/></a:path></a:gradFill>'
        )

    def _paint(self, ctx: _Ctx, which: str, local_box, matrix,
               frame_rot=0.0):
        """DrawingML fill for ``fill`` or ``stroke``, or None for none."""
        style = ctx.style
        value = style.get(which, _INITIAL.get(which))
        alpha = ctx.opacity * _opacity(style.get(f"{which}-opacity"))
        color = _parse_color(value, style.get("color", "black"))
        if color is None:
            return None
        if color[0] == "url":
            xml = self._gradient_xml(color[1], alpha, local_box, matrix,
                                     frame_rot)
            if xml is None and color[2]:
                fallback = _parse_color(color[2], style.get("color", "black"))
                if fallback and fallback[0] != "url":
                    return (f"<a:solidFill>{_color_xml(fallback[0], fallback[1] * alpha)}"
                            "</a:solidFill>")
            return xml
        return f"<a:solidFill>{_color_xml(color[0], color[1] * alpha)}</a:solidFill>"

    def _marker_end(self, ctx: _Ctx, which: str, stroke_width: float):
        value = ctx.style.get(f"marker-{which}")
        color = _parse_color(value) if value else None
        if not color or color[0] != "url":
            return ""
        marker = self._ids.get(color[1])
        if marker is None or marker.tag != "marker":
            return ""
        kind = "triangle"
        shapes = [c for c in marker.iter() if c.tag in (
            "path", "polygon", "polyline", "circle", "ellipse", "rect")]
        if shapes:
            first = shapes[0]
            if first.tag in ("circle", "ellipse"):
                kind = "oval"
            elif first.tag == "rect":
                kind = "diamond"
            else:
                st = self._computed(first, dict(_INITIAL), [marker])
                closed = first.tag == "polygon" or "z" in (
                    first.get("d") or "").lower()
                if not closed or _parse_color(st.get("fill")) is None:
                    kind = "arrow"
        size = _length(marker.get("markerWidth"), default=3.0)
        if marker.get("markerUnits") == "userSpaceOnUse" and stroke_width:
            size /= stroke_width
        grade = "sm" if size < 2.5 else "med" if size < 4 else "lg"
        tag = "a:headEnd" if which == "start" else "a:tailEnd"
        return f'<{tag} type="{kind}" w="{grade}" len="{grade}"/>'

    def _line_xml(self, ctx: _Ctx, local_box, matrix, frame_rot=0.0,
                  markers=False) -> str:
        style = ctx.style
        paint = self._paint(ctx, "stroke", local_box, matrix, frame_rot)
        width_user = _length(style.get("stroke-width", "1"),
                             style.get("font-size", 16.0),
                             math.hypot(*ctx.viewport) / math.sqrt(2), 1.0)
        if paint is None or width_user <= 0:
            return "<a:ln><a:noFill/></a:ln>"
        scale = _mean_scale(matrix)
        width = max(1, min(round(width_user * scale), 20116800))
        cap = {"round": "rnd", "square": "sq"}.get(
            style.get("stroke-linecap", ""), "flat")
        parts = [f'<a:ln w="{width}" cap="{cap}">', paint]
        dashes = [
            _length(v, default=0.0)
            for v in re.split(r"[\s,]+", style.get("stroke-dasharray", "") or "")
            if v and v != "none"
        ]
        if dashes and any(d > 0 for d in dashes):
            if len(dashes) % 2:
                dashes *= 2
            ds = "".join(
                f'<a:ds d="{max(1, round(d / width_user * 100000))}" '
                f'sp="{max(1, round(s / width_user * 100000))}"/>'
                for d, s in zip(dashes[::2], dashes[1::2], strict=True)
            )
            parts.append(f"<a:custDash>{ds}</a:custDash>")
        join = style.get("stroke-linejoin", "miter")
        parts.append({"round": "<a:round/>", "bevel": "<a:bevel/>"}.get(
            join, '<a:miter lim="800000"/>'))
        if markers:
            parts.append(self._marker_end(ctx, "start", width_user))
            parts.append(self._marker_end(ctx, "end", width_user))
        parts.append("</a:ln>")
        return "".join(parts)

    # -- shapes -----------------------------------------------------------

    def _shape(self, el, ctx: _Ctx) -> list:
        tag = el.tag
        fs = ctx.style.get("font-size", 16.0)
        vw, vh = ctx.viewport
        vd = math.hypot(vw, vh) / math.sqrt(2)

        def L(name, ref, default=0.0):  # noqa: N802 — reads like the attribute
            return _length(el.get(name), fs, ref, default)

        if tag == "rect":
            x, y, w, h = L("x", vw), L("y", vh), L("width", vw), L("height", vh)
            if w <= 0 or h <= 0:
                return []
            rx = _length(el.get("rx"), fs, vw, None)
            ry = _length(el.get("ry"), fs, vh, None)
            if rx is None:
                rx = ry or 0.0
            if ry is None:
                ry = rx
            rx, ry = min(rx, w / 2), min(ry, h / 2)
            box = _Box(x, y, w, h)
            if abs(rx - ry) < 1e-6:
                prst = "roundRect" if rx > 0 else "rect"
                adj = round(rx / min(w, h) * 100000) if rx > 0 else None
                made = self._preset(el, ctx, box, prst, adj, "Rectangle")
                if made:
                    return made
            d = (
                f"M{x + rx},{y} H{x + w - rx} A{rx},{ry} 0 0 1 {x + w},{y + ry} "
                f"V{y + h - ry} A{rx},{ry} 0 0 1 {x + w - rx},{y + h} "
                f"H{x + rx} A{rx},{ry} 0 0 1 {x},{y + h - ry} "
                f"V{y + ry} A{rx},{ry} 0 0 1 {x + rx},{y} Z"
            )
            return self._freeform(el, ctx, _parse_path(d), "Rectangle")

        if tag in ("circle", "ellipse"):
            cx, cy = L("cx", vw), L("cy", vh)
            if tag == "circle":
                rx = ry = L("r", vd)
            else:
                rx, ry = L("rx", vw), L("ry", vh)
                if el.get("rx") is None:
                    rx = ry
                if el.get("ry") is None:
                    ry = rx
            if rx <= 0 or ry <= 0:
                return []
            box = _Box(cx - rx, cy - ry, 2 * rx, 2 * ry)
            made = self._preset(el, ctx, box, "ellipse", None, "Oval")
            if made:
                return made
            d = (
                f"M{cx + rx},{cy} A{rx},{ry} 0 0 1 {cx},{cy + ry} "
                f"A{rx},{ry} 0 0 1 {cx - rx},{cy} A{rx},{ry} 0 0 1 {cx},{cy - ry} "
                f"A{rx},{ry} 0 0 1 {cx + rx},{cy} Z"
            )
            return self._freeform(el, ctx, _parse_path(d), "Oval")

        if tag == "line":
            segments = [
                ("M", L("x1", vw), L("y1", vh)),
                ("L", L("x2", vw), L("y2", vh)),
            ]
            return self._freeform(el, ctx, segments, "Line")

        if tag in ("polyline", "polygon"):
            numbers = _numbers(el.get("points"))
            points = list(zip(numbers[::2], numbers[1::2], strict=False))
            if len(points) < 2:
                return []
            segments = [("M",) + points[0]] + [("L",) + p for p in points[1:]]
            if tag == "polygon":
                segments.append(("Z",))
            return self._freeform(el, ctx, segments, "Freeform")

        if tag == "path":
            return self._freeform(el, ctx, _parse_path(el.get("d")), "Freeform")

        return []

    def _preset(self, el, ctx: _Ctx, box: _Box, prst: str, adj, label):
        sx, sy, angle, skewed = _decompose(ctx.matrix)
        if skewed:
            return None
        cx, cy = transform_point(
            ctx.matrix, box.x + box.w / 2, box.y + box.h / 2,
        )
        w, h = box.w * sx, box.h * sy
        out = _Box(cx - w / 2, cy - h / 2, w, h)
        fill = self._paint(ctx, "fill", box, ctx.matrix, angle) or "<a:noFill/>"
        line = self._line_xml(ctx, box, ctx.matrix, angle)
        av = f'<a:gd name="adj" fmla="val {min(adj, 50000)}"/>' if adj else ""
        template = (
            f"<p:sp {_NSDECLS}><p:nvSpPr>"
            '<p:cNvPr id="{id}" name="{name}"/><p:cNvSpPr/><p:nvPr/>'
            f"</p:nvSpPr><p:spPr>{_xfrm(out, angle)}"
            f'<a:prstGeom prst="{prst}"><a:avLst>{av}</a:avLst></a:prstGeom>'
            f"{fill}{line}</p:spPr></p:sp>"
        )
        return [_Shape(out, _brace_safe(template), self._name(el, label))]

    def _freeform(self, el, ctx: _Ctx, segments, label) -> list:
        if not segments:
            return []
        local_box = _segments_bbox(segments) or _Box(0, 0, 0, 0)
        world = _transform_segments(ctx.matrix, segments)
        box = _segments_bbox(world)
        if box is None:
            return []
        closed = any(s[0] == "Z" for s in segments)
        fill = self._paint(ctx, "fill", local_box, ctx.matrix)
        line = self._line_xml(ctx, local_box, ctx.matrix, markers=not closed)

        kinds = [s[0] for s in world]
        if kinds == ["M", "L"]:
            return self._connector(el, world, line, label)

        w, h = max(round(box.w), 1), max(round(box.h), 1)

        def pt(x, y):
            return f'<a:pt x="{round(x - box.x)}" y="{round(y - box.y)}"/>'

        cmds = []
        for seg in world:
            if seg[0] == "M":
                cmds.append(f"<a:moveTo>{pt(seg[1], seg[2])}</a:moveTo>")
            elif seg[0] == "L":
                cmds.append(f"<a:lnTo>{pt(seg[1], seg[2])}</a:lnTo>")
            elif seg[0] == "C":
                cmds.append(
                    "<a:cubicBezTo>" + pt(seg[1], seg[2]) + pt(seg[3], seg[4])
                    + pt(seg[5], seg[6]) + "</a:cubicBezTo>"
                )
            else:
                cmds.append("<a:close/>")
        path_attrs = "" if fill else ' fill="none"'
        template = (
            f"<p:sp {_NSDECLS}><p:nvSpPr>"
            '<p:cNvPr id="{id}" name="{name}"/><p:cNvSpPr/><p:nvPr/>'
            f"</p:nvSpPr><p:spPr>{_xfrm(box)}"
            "<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>"
            '<a:rect l="0" t="0" r="r" b="b"/>'
            f'<a:pathLst><a:path w="{w}" h="{h}"{path_attrs}>'
            f"{''.join(cmds)}</a:path></a:pathLst></a:custGeom>"
            f"{fill or '<a:noFill/>'}{line}</p:spPr></p:sp>"
        )
        return [_Shape(box, _brace_safe(template), self._name(el, label))]

    def _connector(self, el, world, line, label) -> list:
        (_, x1, y1), (_, x2, y2) = world
        box = _Box(min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1))
        template = (
            f"<p:cxnSp {_NSDECLS}><p:nvCxnSpPr>"
            '<p:cNvPr id="{id}" name="{name}"/><p:cNvCxnSpPr/><p:nvPr/>'
            f"</p:nvCxnSpPr><p:spPr>{_xfrm(box, 0, x2 < x1, y2 < y1)}"
            '<a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
            f"{line}</p:spPr></p:cxnSp>"
        )
        return [_Shape(box, _brace_safe(template), self._name(el, label))]

    # -- images -----------------------------------------------------------

    def _image(self, el, ctx: _Ctx) -> list:
        from pptx.parts.image import Image

        href = (el.get("href") or "").strip()
        data = None
        if href.startswith("data:"):
            header, _, payload = href.partition(",")
            try:
                if ";base64" in header:
                    data = base64.b64decode(payload)
                else:
                    from urllib.parse import unquote_to_bytes
                    data = unquote_to_bytes(payload)
            except Exception:
                data = None
        elif href and self._resolve_href is not None:
            data = self._resolve_href(href)
        if not data:
            self._warn("Images that couldn't be read were skipped")
            return []
        try:
            image = Image.from_blob(data)
            px_w, px_h = image.size
            _, r_id = self._slide_part.get_or_add_image_part(io.BytesIO(data))
        except Exception:
            self._warn("Images in formats PowerPoint can't embed were skipped")
            return []

        fs = ctx.style.get("font-size", 16.0)
        vw, vh = ctx.viewport
        x = _length(el.get("x"), fs, vw, 0.0)
        y = _length(el.get("y"), fs, vh, 0.0)
        w = _length(el.get("width"), fs, vw, None)
        h = _length(el.get("height"), fs, vh, None)
        if w is None and h is None:
            w, h = float(px_w), float(px_h)
        elif w is None:
            w = h * px_w / px_h
        elif h is None:
            h = w * px_h / px_w
        if w <= 0 or h <= 0:
            return []
        par = (el.get("preserveAspectRatio") or "xMidYMid meet").split()
        if par and par[0] != "none" and px_w and px_h:
            fit = _viewbox_matrix((0, 0, px_w, px_h), w, h, " ".join(par))
            if "slice" in par:
                self._warn("Cropped images (slice) were fitted instead")
                fit = _viewbox_matrix((0, 0, px_w, px_h), w, h, par[0] + " meet")
            x, y = x + fit.e, y + fit.f
            w, h = px_w * fit.a, px_h * fit.d
        sx, sy, angle, skewed = _decompose(ctx.matrix)
        if skewed:
            self._warn("Skewed images were placed unskewed")
        cx, cy = transform_point(ctx.matrix, x + w / 2, y + h / 2)
        out = _Box(cx - w * sx / 2, cy - h * sy / 2, w * sx, h * sy)
        alpha = ""
        if ctx.opacity < 0.999:
            alpha = f'<a:alphaModFix amt="{round(ctx.opacity * 100000)}"/>'
        template = (
            f"<p:pic {_NSDECLS}><p:nvPicPr>"
            '<p:cNvPr id="{id}" name="{name}"/>'
            '<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr>'
            "<p:nvPr/></p:nvPicPr>"
            f'<p:blipFill><a:blip r:embed="{r_id}">{alpha}</a:blip>'
            "<a:stretch><a:fillRect/></a:stretch></p:blipFill>"
            f"<p:spPr>{_xfrm(out, angle)}"
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
        )
        return [_Shape(out, _brace_safe(template), self._name(el, "Picture"))]

    # -- text -------------------------------------------------------------

    def _text_lines(self, el, ctx: _Ctx) -> list[_Line]:
        lines: list[_Line] = []
        fs = ctx.style.get("font-size", 16.0)
        vw, vh = ctx.viewport
        x = _first_length(el.get("x"), fs, vw, 0.0)
        y = _first_length(el.get("y"), fs, vh, 0.0)
        x += _first_length(el.get("dx"), fs, vw, 0.0)
        y += _first_length(el.get("dy"), fs, vh, 0.0)
        current = _Line(x, y)
        lines.append(current)

        def add(text, style):
            nonlocal current
            if text:
                current.runs.append(_Run(text, style))

        def walk(node, node_ctx):
            nonlocal current
            add(node.text, node_ctx.style)
            for child in node:
                if child.tag not in ("tspan", "a", "textPath"):
                    add(child.tail, node_ctx.style)
                    continue
                child_ctx = self._child_ctx(child, node_ctx)
                if child_ctx.style.get("display") == "none":
                    add(child.tail, node_ctx.style)
                    continue
                cfs = child_ctx.style.get("font-size", fs)
                has_x = child.get("x") is not None
                has_y = child.get("y") is not None
                dx = _first_length(child.get("dx"), cfs, vw, 0.0)
                dy = _first_length(child.get("dy"), cfs, vh, 0.0)
                if has_x or has_y or abs(dy) > 1e-9:
                    nx = _first_length(child.get("x"), cfs, vw, 0.0) if has_x else current.x
                    ny = _first_length(child.get("y"), cfs, vh, 0.0) if has_y else current.y
                    nx += dx
                    ny += dy
                    if current.runs:
                        current = _Line(nx, ny)
                        lines.append(current)
                    else:
                        current.x, current.y = nx, ny
                walk(child, child_ctx)
                add(child.tail, node_ctx.style)

        walk(el, ctx)
        out = []
        for line in lines:
            runs = []
            for run in line.runs:
                text = re.sub(r"\s+", " ", run.text)
                if runs and runs[-1].text.endswith(" ") and text.startswith(" "):
                    text = text[1:]
                if text:
                    runs.append(_Run(text, run.style))
            if runs:
                runs[0].text = runs[0].text.lstrip()
                runs[-1].text = runs[-1].text.rstrip()
                runs = [r for r in runs if r.text]
            if runs:
                line.runs = runs
                out.append(line)
        return out

    def _text(self, el, ctx: _Ctx) -> list:
        lines = self._text_lines(el, ctx)
        if not lines:
            return []
        style = ctx.style
        anchor = style.get("text-anchor", "start")
        baseline = style.get("dominant-baseline") or style.get(
            "alignment-baseline") or "auto"
        # One box when the lines stack evenly under one anchor,
        # which is how multi-line labels are written; otherwise a
        # box per line so each keeps its own position.
        groups = [lines]
        if len(lines) > 1:
            gaps = [b.y - a.y for a, b in zip(lines, lines[1:], strict=False)]
            even = all(abs(g - gaps[0]) <= 0.05 * abs(gaps[0]) for g in gaps)
            same_x = all(abs(ln.x - lines[0].x) < 0.5 for ln in lines)
            if not (even and same_x and gaps[0] > 0):
                groups = [[ln] for ln in lines]
        nodes = []
        for group in groups:
            node = self._text_box(el, ctx, group, anchor, baseline)
            if node:
                nodes.append(node)
        return self._group(el, nodes, "Text") if len(nodes) > 1 else nodes

    def _text_box(self, el, ctx: _Ctx, lines, anchor, baseline):
        sx, sy, angle, _ = _decompose(ctx.matrix)
        if sy <= 0:
            return None
        sizes = [r.style.get("font-size", 16.0) for ln in lines for r in ln.runs]
        fs = max(sizes)
        pitch = (lines[1].y - lines[0].y) if len(lines) > 1 else fs * _LINE_PITCH_EM
        width = max(
            sum(len(r.text) * r.style.get("font-size", 16.0) * (
                0.6 if _is_bold(r.style) else _CHAR_WIDTH_EM) for r in ln.runs)
            for ln in lines
        )
        height = pitch * len(lines)
        x0 = lines[0].x
        left = {"middle": x0 - width / 2, "end": x0 - width}.get(anchor, x0)
        # Place the box so the first line's baseline (or middle,
        # or top) lands where the SVG puts it.
        y0 = lines[0].y
        if baseline in ("middle", "central"):
            top = y0 - pitch / 2
        elif baseline in ("hanging", "text-before-edge", "text-top"):
            top = y0 - (pitch - fs)
        else:
            top = y0 - (pitch - _DESCENT_EM * fs)
        local = _Box(left, top, width, height)
        cx, cy = transform_point(ctx.matrix, left + width / 2, top + height / 2)
        out = _Box(cx - width * sx / 2, cy - height * sy / 2,
                   width * sx, height * sy)
        algn = {"middle": "ctr", "end": "r"}.get(anchor, "l")
        spacing = ""
        if len(lines) > 1:
            pct = pitch / (fs * _LINE_PITCH_EM)
            spacing = (
                f'<a:lnSpc><a:spcPct val="{round(pct * 100000)}"/></a:lnSpc>'
            )
        paragraphs = []
        for line in lines:
            runs = "".join(
                self._run_xml(run, ctx, local, sy) for run in line.runs
            )
            paragraphs.append(
                f'<a:p><a:pPr algn="{algn}">{spacing}</a:pPr>{runs}</a:p>'
            )
        template = (
            f"<p:sp {_NSDECLS}><p:nvSpPr>"
            '<p:cNvPr id="{id}" name="{name}"/><p:cNvSpPr txBox="1"/><p:nvPr/>'
            f"</p:nvSpPr><p:spPr>{_xfrm(out, angle)}"
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
            f'<p:txBody><a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" '
            f'bIns="0" anchor="t" rtlCol="0"><a:noAutofit/></a:bodyPr>'
            f"<a:lstStyle/>{''.join(paragraphs)}</p:txBody></p:sp>"
        )
        label = " ".join(r.text for r in lines[0].runs)[:40]
        name = el.get("id") or f"Text: {label}"
        return _Shape(out, _brace_safe(template), name)

    def _run_xml(self, run: _Run, ctx: _Ctx, local_box, sy) -> str:
        style = run.style
        size_pt = style.get("font-size", 16.0) * sy / _EMU_PER_PT
        attrs = f'lang="en-US" sz="{max(100, round(size_pt * 100))}"'
        if _is_bold(style):
            attrs += ' b="1"'
        if style.get("font-style") in ("italic", "oblique"):
            attrs += ' i="1"'
        decoration = style.get("text-decoration", "") or ""
        if "underline" in decoration:
            attrs += ' u="sng"'
        if "line-through" in decoration:
            attrs += ' strike="sngStrike"'
        run_ctx = _Ctx(ctx.matrix, style, ctx.ancestors, ctx.opacity,
                       ctx.viewport)
        fill = self._paint(run_ctx, "fill", local_box, ctx.matrix)
        if fill and fill.startswith("<a:gradFill"):
            match = re.search(r"<a:srgbClr[^>]*?(?:/>|>.*?</a:srgbClr>)", fill)
            fill = f"<a:solidFill>{match.group(0)}</a:solidFill>" if match else None
        fill = fill or "<a:noFill/>"
        font = escape(_font_name(style.get("font-family", "")), {'"': "&quot;"})
        return (
            f"<a:r><a:rPr {attrs} dirty=\"0\">{fill}"
            f'<a:latin typeface="{font}"/><a:cs typeface="{font}"/></a:rPr>'
            f"<a:t>{escape(run.text)}</a:t></a:r>"
        )


def _is_bold(style: dict) -> bool:
    weight = str(style.get("font-weight", "normal")).strip().lower()
    if weight in ("bold", "bolder"):
        return True
    try:
        return int(weight) >= 600
    except ValueError:
        return False


def _font_name(families: str) -> str:
    """First concrete family; generic names map to a common font."""
    generic = None
    for family in families.split(","):
        name = family.strip().strip("'\"")
        if not name:
            continue
        key = name.lower()
        if key in _GENERIC_FONTS:
            generic = generic or _GENERIC_FONTS[key]
            continue
        return name
    return generic or "Times New Roman"


def _brace_safe(template: str) -> str:
    """Escape braces except the ``{id}`` / ``{name}`` slots."""
    return (
        template.replace("{", "{{").replace("}", "}}")
        .replace("{{id}}", "{id}").replace("{{name}}", "{name}")
    )


def _strip_namespaces(root: ET.Element) -> None:
    for el in root.iter():
        if isinstance(el.tag, str) and "}" in el.tag:
            el.tag = el.tag.split("}", 1)[1]
        for key in list(el.attrib):
            if "}" in key:
                local = key.split("}", 1)[1]
                value = el.attrib.pop(key)
                if local == "space":
                    continue
                el.attrib.setdefault(local, value)
