"""SVG → PowerPoint export — the SVG viewer's "Export as PowerPoint".

Each test converts a small SVG and reads the slide back with
python-pptx, pinning that elements arrive as native shapes (not
pictures) at the right place, size, and style. One SVG pixel is
9525 EMU, so the expected numbers below are pixel counts times
9525.
"""

from __future__ import annotations

import base64
import io

import pytest

from ._helpers import _make_png_bytes, _require_pptx

PX = 9525
_NS = 'xmlns="http://www.w3.org/2000/svg"'


def _convert(body: str, attrs: str = 'width="200" height="100"'):
    _require_pptx()
    from aic_dc.doc_convert.svg_to_pptx import svg_to_pptx

    data, warnings = svg_to_pptx(f"<svg {_NS} {attrs}>{body}</svg>")
    return _open(data), warnings


def _open(data: bytes):
    from pptx import Presentation

    return Presentation(io.BytesIO(data))


def _shapes(prs):
    return list(prs.slides[0].shapes)


def _xml(shape) -> str:
    from lxml import etree

    return etree.tostring(shape._element).decode()


class TestSlide:
    def test_slide_is_the_svg_size(self):
        prs, _ = _convert('<rect width="10" height="10"/>')
        assert prs.slide_width == 200 * PX
        assert prs.slide_height == 100 * PX

    def test_viewbox_alone_sets_the_size(self):
        prs, _ = _convert(
            '<rect width="10" height="10"/>', 'viewBox="0 0 400 300"',
        )
        assert (prs.slide_width, prs.slide_height) == (400 * PX, 300 * PX)

    def test_viewbox_scales_into_width_and_height(self):
        prs, _ = _convert(
            '<rect x="10" y="10" width="20" height="20"/>',
            'width="200" height="200" viewBox="0 0 100 100"',
        )
        [shape] = _shapes(prs)
        assert (shape.left, shape.top) == (20 * PX, 20 * PX)
        assert (shape.width, shape.height) == (40 * PX, 40 * PX)

    def test_oversized_svg_is_scaled_to_the_slide_limit(self):
        prs, _ = _convert(
            '<rect width="10000" height="10000"/>',
            'width="10000" height="5000"',
        )
        assert prs.slide_width == 51206400
        assert prs.slide_height == 51206400 // 2

    def test_invalid_svg_raises(self):
        _require_pptx()
        from aic_dc.doc_convert.svg_to_pptx import SvgExportError, svg_to_pptx

        with pytest.raises(SvgExportError):
            svg_to_pptx("<svg")
        with pytest.raises(SvgExportError):
            svg_to_pptx(f"<svg {_NS}><defs/></svg>")


class TestShapes:
    def test_rect_is_a_preset_rectangle(self):
        from pptx.enum.shapes import MSO_SHAPE

        prs, _ = _convert(
            '<rect id="box" x="10" y="20" width="30" height="40" '
            'fill="#ff0000" stroke="blue" stroke-width="2"/>'
        )
        [shape] = _shapes(prs)
        assert shape.auto_shape_type == MSO_SHAPE.RECTANGLE
        assert shape.name == "box"
        assert (shape.left, shape.top) == (10 * PX, 20 * PX)
        assert (shape.width, shape.height) == (30 * PX, 40 * PX)
        assert str(shape.fill.fore_color.rgb) == "FF0000"
        assert str(shape.line.color.rgb) == "0000FF"
        assert shape.line.width == 2 * PX

    def test_rounded_rect_keeps_its_corner_radius(self):
        from pptx.enum.shapes import MSO_SHAPE

        prs, _ = _convert('<rect width="40" height="20" rx="5"/>')
        [shape] = _shapes(prs)
        assert shape.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE
        assert shape.adjustments[0] == pytest.approx(0.25)

    def test_rotation_becomes_shape_rotation(self):
        prs, _ = _convert(
            '<rect x="10" y="10" width="40" height="20" '
            'transform="rotate(30 30 20)"/>'
        )
        [shape] = _shapes(prs)
        assert shape.rotation == pytest.approx(30)
        assert shape.width == 40 * PX
        assert shape.left + shape.width / 2 == pytest.approx(30 * PX, abs=1)

    def test_circle_is_an_ellipse(self):
        from pptx.enum.shapes import MSO_SHAPE

        prs, _ = _convert('<circle cx="50" cy="50" r="10" fill="green"/>')
        [shape] = _shapes(prs)
        assert shape.auto_shape_type == MSO_SHAPE.OVAL
        assert (shape.left, shape.width) == (40 * PX, 20 * PX)

    def test_skewed_ellipse_falls_back_to_a_freeform(self):
        prs, _ = _convert(
            '<ellipse cx="50" cy="50" rx="20" ry="10" '
            'transform="skewX(30)"/>'
        )
        [shape] = _shapes(prs)
        assert "<a:custGeom>" in _xml(shape)

    def test_path_is_a_freeform_with_curves(self):
        prs, _ = _convert(
            '<path d="M10 10 Q 30 0 50 10 A 10 10 0 0 1 70 10 Z" '
            'fill="none" stroke="black"/>'
        )
        [shape] = _shapes(prs)
        xml = _xml(shape)
        assert "<a:custGeom>" in xml
        assert xml.count("<a:cubicBezTo>") >= 2
        assert "<a:close/>" in xml
        assert 'fill="none"' in xml

    def test_freeform_frame_is_the_tight_curve_bounds(self):
        # The control point sits at y=-10, but the curve itself
        # only rises to y=0 (its midpoint), so the frame starts
        # there rather than at the control point.
        prs, _ = _convert('<path d="M0 10 Q 50 -10 100 10"/>')
        [shape] = _shapes(prs)
        assert shape.top == pytest.approx(0, abs=2)
        assert shape.height == pytest.approx(10 * PX, abs=2)

    def test_line_is_a_connector_with_arrowhead(self):
        prs, _ = _convert(
            '<defs><marker id="m" markerWidth="7"><path d="M0 0 L10 5 L0 10 z"/>'
            '</marker></defs>'
            '<line x1="100" y1="50" x2="20" y2="10" stroke="black" '
            'marker-end="url(#m)"/>'
        )
        [shape] = _shapes(prs)
        xml = _xml(shape)
        assert xml.startswith("<p:cxnSp")
        assert 'flipH="1"' in xml and 'flipV="1"' in xml
        assert '<a:tailEnd type="triangle"' in xml
        assert (shape.left, shape.top) == (20 * PX, 10 * PX)

    def test_dash_array_becomes_a_custom_dash(self):
        prs, _ = _convert(
            '<rect width="10" height="10" stroke="black" stroke-width="2" '
            'stroke-dasharray="4 2"/>'
        )
        xml = _xml(_shapes(prs)[0])
        assert '<a:ds d="200000" sp="100000"/>' in xml

    def test_opacity_becomes_alpha(self):
        prs, _ = _convert(
            '<g opacity="0.5"><rect width="10" height="10" '
            'fill="rgba(0,0,255,0.5)"/></g>'
        )
        assert '<a:alpha val="25000"/>' in _xml(_shapes(prs)[0])

    def test_linear_gradient_becomes_a_gradient_fill(self):
        prs, _ = _convert(
            '<defs><linearGradient id="g" x2="0" y2="1">'
            '<stop offset="0" stop-color="#fff"/>'
            '<stop offset="1" stop-color="#000"/></linearGradient></defs>'
            '<rect width="10" height="10" fill="url(#g)"/>'
        )
        xml = _xml(_shapes(prs)[0])
        assert "<a:gradFill" in xml
        assert '<a:lin ang="5400000"' in xml

    def test_clip_paths_are_reported(self):
        _, warnings = _convert(
            '<defs><clipPath id="c"><rect width="5" height="5"/></clipPath></defs>'
            '<rect width="10" height="10" clip-path="url(#c)"/>'
        )
        assert warnings == ["Clip paths were ignored"]


class TestStyles:
    def test_stylesheet_classes_apply(self):
        prs, _ = _convert(
            "<style>.box { fill: #123456 } g.hot > rect { stroke: red }</style>"
            '<g class="hot"><rect class="box" width="10" height="10"/>'
            '<rect x="20" width="10" height="10"/></g>'
        )
        [group] = _shapes(prs)
        first, second = group.shapes
        assert str(first.fill.fore_color.rgb) == "123456"
        assert str(first.line.color.rgb) == "FF0000"
        assert str(second.fill.fore_color.rgb) == "000000"

    def test_inline_style_beats_stylesheet_and_attributes(self):
        prs, _ = _convert(
            "<style>rect { fill: blue }</style>"
            '<rect width="10" height="10" fill="green" style="fill: orange"/>'
        )
        assert str(_shapes(prs)[0].fill.fore_color.rgb) == "FFA500"

    def test_fill_none_inherits(self):
        prs, _ = _convert(
            '<g fill="none" stroke="black"><rect width="10" height="10"/></g>'
        )
        assert "<a:noFill/>" in _xml(_shapes(prs)[0])


class TestGroups:
    def test_groups_nest(self):
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        prs, _ = _convert(
            '<g id="outer"><rect width="10" height="10"/>'
            '<g><rect x="20" width="10" height="10"/>'
            '<rect x="40" width="10" height="10"/></g></g>'
        )
        [outer] = _shapes(prs)
        assert outer.shape_type == MSO_SHAPE_TYPE.GROUP
        assert outer.name == "outer"
        kinds = [s.shape_type for s in outer.shapes]
        assert kinds[1] == MSO_SHAPE_TYPE.GROUP

    def test_single_child_groups_flatten(self):
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        prs, _ = _convert('<g><g><rect width="10" height="10"/></g></g>')
        [shape] = _shapes(prs)
        assert shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE

    def test_use_of_a_symbol_maps_its_viewbox(self):
        prs, _ = _convert(
            '<defs><symbol id="s" viewBox="0 0 10 10">'
            '<rect width="10" height="10"/></symbol></defs>'
            '<use href="#s" x="50" y="20" width="30" height="30"/>'
        )
        [shape] = _shapes(prs)
        assert (shape.left, shape.top) == (50 * PX, 20 * PX)
        assert shape.width == 30 * PX


class TestText:
    def test_text_becomes_a_text_box(self):
        prs, _ = _convert(
            '<text x="10" y="50" font-family="Segoe UI, sans-serif" '
            'font-size="16" fill="#336699">Hello &amp; welcome</text>'
        )
        [shape] = _shapes(prs)
        assert shape.has_text_frame
        assert shape.text_frame.text == "Hello & welcome"
        [run] = shape.text_frame.paragraphs[0].runs
        assert run.font.size.pt == pytest.approx(12)
        assert run.font.name == "Segoe UI"
        assert str(run.font.color.rgb) == "336699"
        assert shape.left == 10 * PX

    def test_tspans_become_runs_and_lines(self):
        prs, _ = _convert(
            '<text x="100" y="20" text-anchor="middle">'
            '<tspan x="100" dy="0">Top <tspan font-weight="bold">bold</tspan></tspan>'
            '<tspan x="100" dy="20">Second</tspan></text>'
        )
        [shape] = _shapes(prs)
        first, second = shape.text_frame.paragraphs
        assert [r.text for r in first.runs] == ["Top ", "bold"]
        assert first.runs[1].font.bold
        assert second.runs[0].text == "Second"
        assert 'algn="ctr"' in _xml(shape)
        assert shape.left + shape.width / 2 == pytest.approx(100 * PX, abs=1)

    def test_unevenly_spaced_lines_get_their_own_boxes(self):
        prs, _ = _convert(
            '<text><tspan x="10" y="10">a</tspan><tspan x="60" y="15">b</tspan>'
            '</text>'
        )
        [group] = _shapes(prs)
        assert [s.text_frame.text for s in group.shapes] == ["a", "b"]

    def test_generic_family_maps_to_a_real_font(self):
        prs, _ = _convert('<text font-family="monospace" y="20">x</text>')
        run = _shapes(prs)[0].text_frame.paragraphs[0].runs[0]
        assert run.font.name == "Courier New"


class TestImages:
    def test_data_uri_image_is_a_picture(self):
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        png = base64.b64encode(_make_png_bytes()).decode()
        prs, _ = _convert(
            f'<image x="10" y="10" width="40" height="40" '
            f'href="data:image/png;base64,{png}"/>'
        )
        [shape] = _shapes(prs)
        assert shape.shape_type == MSO_SHAPE_TYPE.PICTURE

    def test_unreadable_image_is_reported(self):
        _, warnings = _convert(
            '<rect width="5" height="5"/><image href="missing.png" '
            'width="10" height="10"/>'
        )
        assert warnings == ["Images that couldn't be read were skipped"]


class TestService:
    def test_export_returns_a_named_pptx(self, doc_convert):
        _require_pptx()
        result = doc_convert.export_svg_to_pptx(
            f'<svg {_NS} width="10" height="10"><rect width="5" height="5"/></svg>',
            "docs/diagram.svg",
        )
        assert result["filename"] == "diagram.pptx"
        assert result["warnings"] == []
        prs = _open(base64.b64decode(result["data"]))
        assert len(_shapes(prs)) == 1

    def test_relative_image_href_reads_from_the_repo(
        self, doc_convert, scan_root
    ):
        _require_pptx()
        (scan_root / "docs").mkdir()
        (scan_root / "docs" / "pic.png").write_bytes(_make_png_bytes())
        result = doc_convert.export_svg_to_pptx(
            f'<svg {_NS} width="10" height="10">'
            '<image href="pic.png" width="5" height="5"/></svg>',
            "docs/diagram.svg",
        )
        assert result["warnings"] == []

    def test_href_outside_the_repo_is_not_read(
        self, doc_convert, scan_root
    ):
        _require_pptx()
        outside = scan_root.parent / "secret.png"
        outside.write_bytes(_make_png_bytes())
        result = doc_convert.export_svg_to_pptx(
            f'<svg {_NS} width="10" height="10"><rect width="5" height="5"/>'
            '<image href="../secret.png" width="5" height="5"/></svg>',
            "diagram.svg",
        )
        assert result["warnings"] == [
            "Images that couldn't be read were skipped"
        ]

    def test_bad_svg_is_an_error_result(self, doc_convert):
        _require_pptx()
        result = doc_convert.export_svg_to_pptx("not svg", "x.svg")
        assert "error" in result

    def test_missing_python_pptx_is_an_error_result(
        self, doc_convert, monkeypatch
    ):
        monkeypatch.setattr(
            type(doc_convert), "_probe_import", staticmethod(lambda name: False),
        )
        result = doc_convert.export_svg_to_pptx("<svg/>", "x.svg")
        assert "python-pptx" in result["error"]
